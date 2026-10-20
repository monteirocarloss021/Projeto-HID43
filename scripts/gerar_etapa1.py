#!/usr/bin/env python3
"""Gera memória, CSV e Word a partir de premissas explícitas; não dimensiona circuitos."""
import csv
import json
from decimal import Decimal, ROUND_CEILING, ROUND_FLOOR
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
D = Decimal


def br(value, places=2):
    return f"{D(value):,.{places}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def luz(area):
    area = D(area)
    return D(100) + D(60) * max(D(0), ((area - 6) / 4).to_integral_value(rounding=ROUND_FLOOR))


def tug(n):
    return D(600) * min(3, n) + D(100) * max(0, n - 3)


def calcular(data):
    rooms = data["ambientes"]
    loads = data["equipamentos"]
    for room in rooms:
        room["iluminacao_va"] = luz(room["area_m2"])
        room["tug_va"] = tug(room["tug"])
    for item in loads:
        p, u, fp = (D(item[k]) for k in ("potencia_w", "tensao_v", "fp"))
        if not (p > 0 and u in (127, 220) and 0 < fp <= 1):
            raise ValueError(f"Premissa elétrica inválida: {item['ids']}")
        item["corrente_a"] = p / (u * fp)
        item["aparente_va"] = p / fp
        item["acima_10a"] = item["corrente_a"] > 10
    totals = {
        "pontos_luz": sum(r["pontos_luz"] for r in rooms),
        "iluminacao_va": sum(r["iluminacao_va"] for r in rooms),
        "pontos_tug": sum(r["tug"] for r in rooms),
        "tug_va": sum(r["tug_va"] for r in rooms),
        "pontos_especificos": sum(len(e["ids"]) for e in loads),
        "especificos_w": sum(D(e["potencia_w"]) * len(e["ids"]) for e in loads),
        "especificos_va": sum(e["aparente_va"] * len(e["ids"]) for e in loads),
    }
    totals["soma_aparente_va"] = totals["iluminacao_va"] + totals["tug_va"] + totals["especificos_va"]
    return rooms, loads, totals


def verificar(rooms, loads, totals):
    """Valores de controle independentes da composição das tabelas."""
    assert [r["iluminacao_va"] for r in rooms] == [340, 280, 220, 100, 100]
    assert [r["tug_va"] for r in rooms] == [1900, 1900, 2000, 1900, 1200]
    for room in rooms:
        if room["id"] != "S":
            passo = D(5) if room["id"] == "V" else D("3.5")
            assert room["tug"] == int((D(room["perimetro_m"]) / passo).to_integral_value(rounding=ROUND_CEILING))
    assert totals["iluminacao_va"] == 1040
    assert totals["tug_va"] == 8900
    assert totals["pontos_tug"] == 19
    assert totals["pontos_especificos"] == 14
    assert totals["especificos_w"] == 37860
    assert abs(totals["soma_aparente_va"] - D("48422.89473684210526315789474")) < D("0.000001")
    above = [i for e in loads if e["acima_10a"] for i in e["ids"]]
    assert above == ["CH1", "CH2", "CH3", "CH4", "CH5", "CH6", "ML1"]
    assert luz("10.01") == 160
    assert luz("8.56") + luz("1.45") == 200
    assert tug(6) == 2100
    assert tug(4) + tug(2) == 3100
    assert (D(1200) / 127) < 10 < (D(1200) / (127 * D("0.8")))
    assert D(1500) / 220 < D(1500) / (220 * D("0.95"))
    assert len({i for e in loads for i in e["ids"]}) == 14
    return above


class Report:
    def __init__(self):
        self.doc = Document()
        self.md = []
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.08
        for name in ("Title", "Heading 1", "Heading 2"):
            self.doc.styles[name].font.name = "Calibri"
            self.doc.styles[name].font.color.rgb = RGBColor.from_string("163B50")
        self.doc.styles["Title"].font.size = Pt(25)
        self.doc.styles["Heading 1"].font.size = Pt(17)
        self.doc.styles["Heading 2"].font.size = Pt(12)
        self.configure(self.doc.sections[0], False)
        header = self.doc.sections[0].header.paragraphs[0]
        header.text = "HID-43  /  INSTALAÇÕES ELÉTRICAS                         ETAPA 1"
        header.style = self.doc.styles["Caption"]
        footer = self.doc.sections[0].footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Carlos A. Monteiro da Cunha  •  ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

    def configure(self, section, landscape):
        section.page_width = Cm(29.7 if landscape else 21)
        section.page_height = Cm(21 if landscape else 29.7)
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
        section.header_distance = Cm(0.65)
        section.footer_distance = Cm(0.65)

    def page(self, landscape=False, section=False):
        if section:
            self.configure(self.doc.add_section(WD_SECTION_START.NEW_PAGE), landscape)
        else:
            self.doc.add_page_break()

    def title(self, text, level=1):
        self.doc.add_heading(text, level)
        self.md.append(f"{'#' * (level + 1)} {text}\n")

    def p(self, text, bold=False):
        p = self.doc.add_paragraph()
        p.add_run(text).bold = bold
        self.md.append(text + "\n")
        return p

    def table(self, headers, rows, widths=None, font=9):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"
        table.autofit = False
        for i, label in enumerate(headers):
            table.rows[0].cells[i].text = str(label)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
        if widths:
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        for row in table.rows:
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.line_spacing = 1
                    for run in p.runs:
                        run.font.size = Pt(font)
        self.md.append("| " + " | ".join(headers) + " |")
        self.md.append("| " + " | ".join("---" for _ in headers) + " |")
        for row in rows:
            self.md.append("| " + " | ".join(str(x).replace("\n", "<br>") for x in row) + " |")
        self.md.append("")
        return table


def main():
    data = json.loads((ROOT / "etapa-1/dados.json").read_text())
    rooms, loads, totals = calcular(data)
    above = verificar(rooms, loads, totals)
    report = Report()
    report.title("Projeto HID-43\nEtapa 1 · Previsão de cargas", 0)
    report.p(data["autor"], True)
    report.p("Trabalho individual • Sistema 220/127 V • Revisão: 20/10/2026")
    report.p("Entrega indicada no enunciado: 27/10/2026, às 23h59. O canal deve ser confirmado no Classroom, pois o enunciado de E1 contém o campo “[CANAL DE ENTREGA]”.")
    report.title("1. Base de projeto e continuidade de E0")
    report.p("Este relatório desenvolve a previsão de iluminação, tomadas de uso geral (TUG) e pontos de uso específico para o vestiário, o sanitário, a cozinha e a área de serviço com DML da planta AM008.24-01.A001. Os recintos de serviço e DML são calculados separadamente e somados no resultado do conjunto.")
    report.p("Mantém-se o enquadramento adotado em E0: ambientes coletivos de apoio, com critérios da seção 9.5 da ABNT NBR 5410:2004 empregados POR ANALOGIA. A existência de alojamento na edificação não foi usada para estender automaticamente essa seção a todos os setores. A seção 9.1 permanece pertinente ao vestiário por conter chuveiros.")
    report.title("O que foi conferido na planta", 2)
    report.p("Áreas de E0 preservadas: vestiário 23,98 m²; sanitário 19,84 m²; cozinha 16,94 m²; serviço 8,56 m² e DML 1,45 m². Foram identificados seis boxes de banho, quatro lavatórios no sanitário, duas cubas na cozinha e um tanque no DML.")
    report.p("A cota da área de serviço é 2,75 m. Adota-se P de referência = 2 × (4,00 + 2,75) = 13,50 m, da envolvente conservadora, sem descontar recortes de DML e depósito de gás; não é uma medição do perímetro útil. A estimativa de 13,00 m de E0 também resulta em quatro TUG. Os dois símbolos circulares junto à porta PAVZ 08 pertencem ao depósito de gás, não a máquinas de lavar.")
    report.title("Hipóteses e alcance", 2)
    report.p("Adota-se, para esta previsão acadêmica, aquecimento elétrico nas seis duchas, uma máquina de lavar e os aparelhos de apoio discriminados nas páginas seguintes. Potências, tensões escolhidas e fatores de potência são estimativas declaradas, não dados medidos nem valores de catálogo. O fogão é considerado a gás, sem carga elétrica, com ignição manual; a planta indica depósito de gás, mas o equipamento não está especificado.")
    report.p("Referências conferidas no arquivo NBR-5410.pdf recebido: ABNT NBR 5410:2004, segunda edição, versão corrigida em 17/03/2008. Também foram consultados os PDFs das aulas 28, 29, 30 e 31. As páginas normativas utilizadas estão identificadas ao final; não se atribuem à norma os valores estimados dos equipamentos.", True)

    report.page(landscape=True, section=True)
    report.title("2. Quadro de previsão de cargas")
    report.p("P = potência elétrica ativa de entrada por equipamento; FP = fator de potência global adotado; I = corrente por equipamento. Iluminação e TUG são totais por recinto, lançados apenas na primeira linha. “—” nas demais linhas evita repetir cargas.")
    rows = []
    for room in rooms:
        room_loads = [e for e in loads if e["ambiente"] == room["id"]]
        first = True
        if not room_loads:
            rows.append([room["nome"], "1 / 100", "2 / 1.200", "Sem equipamento específico", "—", "—", "—", "—", "L; T-b; P-a"])
        for e in room_loads:
            for equipment_id in e["ids"]:
                refs = ("L; T-" + ("e" if room["id"] == "V" else "a" if room["id"] == "S" else "b") + "; P-a; E") if first else "E"
                if e["acima_10a"]:
                    refs += "; C"
                if equipment_id.startswith("CH"):
                    refs += "; H"
                rows.append([
                    room["nome"] if first else "↳",
                    f"{room['pontos_luz']} / {br(room['iluminacao_va'],0)}" if first else "—",
                    f"{room['tug']} / {br(room['tug_va'],0)}" if first else "—",
                    f"{equipment_id} · {e['nome']}",
                    br(e["potencia_w"], 0), e["tensao_v"], br(e["fp"]),
                    br(e["corrente_a"]) + (" *" if e["acima_10a"] else ""), refs])
                first = False
    report.table(["Ambiente", "Luz\npontos / VA", "TUG\npontos / VA", "Ponto de uso específico", "P (W)", "U (V)", "FP", "I (A)", "Itens NBR¹"], rows,
                 [2.35, 2.05, 2.05, 5.6, 1.55, 1.4, 1.3, 1.7, 4.3], 8.4)
    report.p("¹ L = 9.5.2.1.1 e 9.5.2.1.2; T-a/b/e = alíneas de 9.5.2.2.1; P-a = 9.5.2.2.2(a); E = 4.2.1.2.1 e 4.2.1.2.3; C = 9.5.3.1; H = 9.5.2.3. Seção 9.5 utilizada por analogia. * Corrente calculada superior a 10 A.")
    report.p(f"Totais: 5 pontos de luz / 1.040 VA; 19 pontos de TUG / 8.900 VA; 14 pontos específicos / {br(totals['especificos_w'],0)} W. Área de serviço + DML: 2 pontos de luz / 200 VA e 6 TUG / 3.100 VA.", True)
    report.p("As tomadas dedicadas aos equipamentos não integram as 19 TUG. “Ponto específico” não significa automaticamente circuito exclusivo; esse critério será aplicado na Etapa 2.")

    report.page(section=True)
    report.title("3. Memória de iluminação e TUG")
    report.title("3.1 Iluminação — 9.5.2.1.1 e 9.5.2.1.2", 2)
    report.p("Por analogia, adota-se no mínimo um ponto fixo de luz no teto, comandado por interruptor, em cada recinto. Para A ≤ 6 m²: 100 VA. Para A > 6 m²: 100 VA + 60 VA por bloco INTEIRO adicional de 4 m². Fórmula: S = 100 + 60 × piso[(A − 6)/4]. Frações do bloco adicional não são arredondadas para cima.")
    light_rows = []
    for room in rooms:
        a = D(room["area_m2"])
        blocks = max(0, int(((a-6)/4).to_integral_value(rounding=ROUND_FLOOR)))
        calc = "100 VA (A ≤ 6)" if a <= 6 else f"piso[({br(a)} − 6)/4] = {blocks}; 100 + 60 × {blocks}"
        light_rows.append([room["nome"], br(a), calc, f"{br(room['iluminacao_va'],0)} VA"])
    report.table(["Recinto", "A (m²)", "Cálculo", "Carga"], light_rows, [3.1, 1.8, 9.9, 2.8])
    report.p("Total: 340 + 280 + 220 + 100 + 100 = 1.040 VA. São cinco pontos mínimos de luz. Esse valor é previsão de carga, não potência real de lâmpadas nem demonstração de iluminância; quantidade final, posição e luminárias exigem estudo próprio (4.2.1.2.2).")
    report.title("3.2 Quantidade de TUG — 9.5.2.2.1", 2)
    tug_rows = [
        ["Vestiário", "e, por analogia", "teto(20,00/5) = 4", "4"],
        ["Sanitário", "a + critério de uso", "1 ponto para cada um dos 4 lavatórios", "4"],
        ["Cozinha", "b, por analogia", "teto(16,50/3,50) = 5", "5"],
        ["Área de serviço", "b, por analogia", "teto(13,50/3,50) = 4", "4"],
        ["DML", "b, local análogo", "teto(4,80/3,50) = 2", "2"]]
    report.table(["Recinto", "Alínea adotada", "Cálculo / decisão", "Pontos"], tug_rows, [3.1, 4.0, 8.0, 2.5])
    report.p("Vestiário: sem lavatório, usa-se o critério de demais dependências com área superior a 6 m² para a quantidade; isso não afasta a seção 9.1. Sanitário: quatro pontos são escolha para uso coletivo, acima do mínimo de um ponto junto ao lavatório; a alínea (a) não exige uma tomada por cuba. Cozinha: ao menos dois dos cinco pontos sobre a bancada da pia.")
    report.title("3.3 Carga das TUG — 9.5.2.2.2(a)", 2)
    report.p("Adotam-se 600 VA por ponto até os três primeiros e 100 VA nos excedentes, em cada recinto separadamente. No vestiário, adota-se o critério de carga de banheiro por conter banho, de forma conservadora. No DML, o tanque fundamenta a analogia a área de serviço. Não se utiliza a redução facultativa para dois pontos de 600 VA.")
    report.table(["Recinto", "Memória (VA)", "Total (VA)"], [
        [r["nome"], f"{min(3,r['tug'])} × 600 + {max(0,r['tug']-3)} × 100", br(r["tug_va"],0)]
        for r in rooms], [4.6, 8.5, 4.5], 8.8)
    report.p("Total de TUG: 1.900 + 1.900 + 2.000 + 1.900 + 1.200 = 8.900 VA.")

    report.page()
    report.title("4. Equipamentos e origem das hipóteses")
    report.p("Os aparelhos abaixo compõem um cenário de previsão, não uma especificação de compra. A norma não fornece estas potências ou fatores de potência. A referência geral é a previsão dos equipamentos e condições de utilização (4.2.1.2.1 e 4.2.1.2.3); cada valor adotado tem a origem indicada.")
    assumption_rows = []
    for e in loads:
        ids = "CH1–CH6" if len(e["ids"]) > 1 else e["ids"][0]
        assumption_rows.append([ids + "\n" + e["nome"], f"{br(e['potencia_w'],0)} W\n{e['tensao_v']} V\nFP {br(e['fp'])}", e["hipotese"]])
    report.table(["Ponto / equipamento", "Valores por unidade", "Origem e decisão de projeto"], assumption_rows, [4.2, 3.0, 10.4], 9)
    report.title("Critérios de escolha e não duplicação", 2)
    report.p("220 V é adotado para chuveiros e alguns aparelhos de maior potência; 127 V para os demais modelos hipotéticos. Considera-se ligação monofásica entre fase e neutro em 127 V ou entre duas fases em 220 V. Não se aplica √3 às correntes desses equipamentos individuais. A distinção entre tensão de linha e de fase é apresentada na aula 31, páginas 5–9.")
    report.p("Geladeira, máquina, ventiladores e exaustores têm motor; FP 0,80 é uma hipótese de cálculo, não constante universal. Secador e micro-ondas usam FP global 0,95 estimado. Todas as potências são de entrada elétrica; não se usam potência mecânica útil nem potência de cocção.")
    report.p("Adota-se uma máquina de lavar como hipótese de uso, sem posição executiva definida. Os dois símbolos circulares junto à porta PAVZ 08 são do depósito de gás e não fundamentam qualquer contagem de máquinas. A posição da lavadora, sua ligação hidráulica e a compatibilidade com o depósito deverão ser verificadas na etapa de locação. No DML, o tanque não é carga elétrica; não se incluem secadora, forno elétrico ou aquecimento de pia sem previsão específica.")
    report.p("As seis duchas elétricas dominam a carga. Se o aquecimento for central ou a gás, as seis cargas de 5.500 W deverão ser retiradas ou substituídas pelos equipamentos efetivos, antes de dimensionar a instalação. Modelos, placas e manuais substituirão as estimativas na continuação do projeto.")

    report.page()
    report.title("5. Correntes dos pontos específicos")
    report.p("Para os equipamentos monofásicos considerados: P = U × I × FP; portanto I = P/(U × FP). A potência aparente por unidade é S = P/FP. O FP é o fator de potência global, inclusive quando existirem componentes eletrônicos. Arredonda-se apenas a apresentação final para duas casas.")
    current_rows = []
    for e in loads:
        ids = "CH1–CH6 (cada)" if len(e["ids"]) > 1 else e["ids"][0]
        current_rows.append([ids, f"{br(e['potencia_w'],0)} / ({e['tensao_v']} × {br(e['fp'])})", br(e["corrente_a"]), "Sim" if e["acima_10a"] else "Não"])
    report.table(["Ponto", "I = P/(U × FP)", "I (A)", "I > 10 A?"], current_rows, [4.0, 8.0, 2.8, 2.8], 9.5)
    report.p("Acima de 10 A: CH1, CH2, CH3, CH4, CH5 e CH6 (25,00 A cada) e ML1 (11,81 A). São sete pontos a reservar para circuitos independentes na Etapa 2, por adoção de 9.5.3.1 por analogia. Essa previsão não define ainda cabos, disjuntores, DR ou balanceamento de fases.", True)
    report.p("Os demais pontos ficam abaixo de 10 A neste cenário, mas isso não autoriza agrupá-los sem analisar uso, fabricante e demais critérios de divisão. Correntes de partida, regime e fator de potência real dos motores não foram usados para dimensionamento nesta etapa.")
    report.title("5.1 Consolidação de cargas", 2)
    report.table(["Grandeza", "Resultado"], [
        ["Iluminação mínima prevista", "1.040 VA"],
        ["TUG previstas", "8.900 VA"],
        ["Soma das potências ativas dos pontos específicos", f"{br(totals['especificos_w'],0)} W"],
        ["Soma escalar de P/FP dos pontos específicos", f"{br(totals['especificos_va'])} VA"],
        ["Soma escalar das cargas aparentes previstas", f"{br(totals['soma_aparente_va'])} VA"]], [11.2, 6.4])
    report.p("A soma escalar é um indicador conservador da carga instalada prevista, sem fator de demanda. Não é demanda calculada, não representa uma soma vetorial das potências e não deve ser dividida por uma tensão única para escolher alimentador. Não se somam diretamente os 37.860 W com os valores em VA de iluminação e TUG.")
    report.title("6. Ligação dos chuveiros — 9.5.2.3")
    report.p("Cada chuveiro é ligado diretamente ao ponto de utilização, por conexão apropriada à corrente e às instruções do fabricante, sem tomada e plugue. No cenário escolhido, são seis ligações diretas em 220 V. O aterramento, a equipotencialização, as restrições dos volumes da seção 9.1 e o DR de alta sensibilidade, IΔn ≤ 30 mA, previsto em E0 por 5.1.3.2.2(a), devem ser atendidos.")
    report.p("A localização das TUG e equipamentos do vestiário deverá ser verificada em relação aos volumes de proteção (9.1.2.1 e requisitos associados da seção 9.1). A contagem por perímetro não autoriza instalar tomadas nos boxes nem em posições proibidas.")

    report.page()
    report.title("7. Registro da IA e revisão")
    report.title("Ponto vago corrigido — símbolos na área de serviço", 2)
    report.p("A redação inicial ainda tratava genericamente as posições junto à área de serviço como equipamentos não confirmados, prolongando a ambiguidade de E0. Ao ampliar a A001, foi lida a identificação “DEPÓSITO DE GÁS” ao lado dos dois símbolos circulares. A referência a possíveis máquinas foi corrigida: esses símbolos não representam lavadoras. ML1 é apenas uma hipótese de uso adotada nesta etapa, não um aparelho extraído da planta.")
    report.title("Pergunta de controle 1 — serviço e DML", 2)
    report.p("Teste aplicado à solução da IA: “Posso somar as áreas de serviço e DML antes de calcular a iluminação e aplicar uma única sequência de 600/100 VA às seis TUG?” A comparação numérica mostra o problema: tratar 10,01 m² como um só recinto resultaria em 160 VA; calcular 8,56 m² e 1,45 m² separadamente resulta em 100 + 100 = 200 VA.")
    report.p("Nas TUG, uma única sequência para seis pontos daria 3 × 600 + 3 × 100 = 2.100 VA. A apuração separada resulta em 1.900 + 1.200 = 3.100 VA. A hipótese de agrupar foi rejeitada porque a planta mostra recintos distintos. A regra de atribuição por ambiente de 9.5.2.2.2(a), usada por analogia, foi conferida na página impressa 184 (página 192 do PDF).")
    report.title("Pergunta de controle 2 — fator de potência", 2)
    report.p("Teste aplicado à IA: “Uma máquina de 1.200 W em 127 V fica abaixo de 10 A se o FP for 0,80?” Usar apenas P/U daria 9,45 A e uma conclusão incorreta. A expressão completa produz 1.200/(127 × 0,80) = 11,81 A. O ponto ML1 foi, portanto, marcado acima de 10 A. Isso mostra por que não se trata motor como carga resistiva.")
    report.p("Este registro documenta testes de consistência efetivamente realizados. A IA foi usada para organizar e conferir a solução; a fonte normativa foi o PDF fornecido, e não a resposta da IA.")
    report.title("8. Rastreabilidade e limites da revisão")
    report.table(["Verificação", "Situação"], [
        ["Áreas, seis duchas e recintos na A001", "Conferidos visualmente na planta recebida."],
        ["Memórias, unidades, somas e limiar de 10 A", "Recalculados por script e comparados com valores de controle."],
        ["Compatibilidade com E0", "Mantida a analogia à seção 9.5; registrada a revisão do perímetro de serviço."],
        ["Potências, modelos e fatores de potência reais", "Não disponíveis; substituídos por estimativas explícitas de projeto."],
        ["NBR 5410 e materiais das aulas", "Itens conferidos no PDF recebido e conceitos confrontados com os slides."],
        ["Circuitos, cabos, disjuntores e posições definitivas", "Fora da previsão de cargas; desenvolver nas próximas etapas."]], [8.0, 9.6])
    report.title("Relação com as aulas", 2)
    report.p("Aula 28, páginas 5–9: motores, fator de potência e distinção entre potência útil e absorvida. Aula 31, páginas 5–9 e 34: tensões de linha/fase e FP = P/S. Aula 29, páginas 4–6 e 24: condutores, comando e representação, a desenvolver nas etapas seguintes. Aula 30, página 17: circuitos exclusivos para cargas acima de 10 A, confrontado com 9.5.3.1.")

    report.page()
    report.title("9. Referências e conferência normativa")
    report.p("No arquivo recebido, a página do PDF corresponde à página impressa acrescida de oito. As referências abaixo foram lidas diretamente no documento normativo, sem misturar a edição de 2004 com a data da versão corrigida de 2008.")
    report.table(["Decisão / assunto", "Item da NBR 5410:2004", "Pág. impressa / PDF"], [
        ["Potência absorvida e pontos específicos", "4.2.1.2.1 e 4.2.1.2.3", "13 / 21"],
        ["Critério geral de iluminação", "4.2.1.2.2", "13 / 21"],
        ["DR de alta sensibilidade", "5.1.3.2.2(a) e (e), notas", "49–50 / 57–58"],
        ["Local de banho e volumes de proteção", "9.1.1 e 9.1.2.1", "170–172 / 178–180"],
        ["Campo de aplicação; um ponto fixo de luz", "9.5.1 e 9.5.2.1.1", "182–183 / 190–191"],
        ["Carga mínima de iluminação por área", "9.5.2.1.2(a) e (b)", "183 / 191"],
        ["Quantidade de TUG e alíneas adotadas", "9.5.2.2.1(a), (b) e (e)", "183 / 191"],
        ["600 VA nos três primeiros; 100 VA excedentes", "9.5.2.2.2(a)", "184 / 192"],
        ["Ligação direta do aquecedor de água", "9.5.2.3", "184 / 192"],
        ["Ponto dedicado com corrente superior a 10 A", "9.5.3.1", "184 / 192"]], [7.0, 6.0, 4.6], 9)
    report.title("Documentos utilizados", 2)
    report.p("[1] E0_Monteiro.docx — levantamento individual recebido, fonte das áreas e dos critérios adotados.")
    report.p("[2] AM008.24-01.A001 — Plantas baixas de execução (pavimento térreo e sala de observação), PDF recebido, folha única.")
    report.p("[3] Projeto HID-43 — Etapa 0: Levantamento — enunciado recebido.")
    report.p("[4] Projeto HID-43 — Etapa 1: Previsão de cargas — enunciado recebido, entrega em 27/10/2026.")
    report.p("[5] ABNT. NBR 5410:2004 — Instalações elétricas de baixa tensão. Segunda edição, 30/09/2004; versão corrigida, 17/03/2008. Arquivo NBR-5410.pdf recebido, 217 páginas digitais, 209 páginas numeradas.")
    report.p("[6] PIMENTEL, Marcio A. S. HID-43 — Instalações Elétricas. PDFs das aulas 28 (25 páginas), 29 (35 páginas), 30 (22 páginas) e 31 (65 páginas), recebidos como material de apoio da disciplina. Foram consultados os slides, sem alegação de reprodução dos vídeos.")
    report.p("Não foram utilizados catálogos como fonte de potência. Os valores estimados são hipóteses autorais desta previsão, apresentadas para permitir revisão e substituição sem alterar o método.")

    out = ROOT / "entregas"
    out.mkdir(exist_ok=True)
    report.doc.core_properties.title = "Projeto HID-43 — Etapa 1: Previsão de cargas"
    report.doc.core_properties.author = data["autor"]
    report.doc.core_properties.subject = "Previsão acadêmica de cargas; premissas explícitas e memória de cálculo"
    report.doc.save(out / "E1_Monteiro.docx")
    (ROOT / "etapa-1/relatorio.md").write_text("\n".join(report.md), encoding="utf-8")
    with (ROOT / "etapa-1/quadro_cargas.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";", lineterminator="\n")
        writer.writerow(["Ambiente", "Luz_pontos_VA", "TUG_pontos_VA", "Equipamento", "P_W", "U_V", "FP", "I_A", "Referencias"])
        writer.writerows(rows)
    results = {
        "totais": totals,
        "pontos_acima_10a": above,
        "premissas": "Dados de potência e FP estimados; itens conferidos no NBR-5410.pdf recebido, versão corrigida de 17/03/2008.",
        "equipamentos": [{"ids": e["ids"], "corrente_a": e["corrente_a"], "aparente_va_por_unidade": e["aparente_va"]} for e in loads]
    }
    (ROOT / "etapa-1/resultados.json").write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str) + "\n")
    print(json.dumps(results, ensure_ascii=False, indent=2, default=str))
    print("Validação dos cálculos: OK. Word, Markdown, CSV e JSON gerados.")


if __name__ == "__main__":
    main()
